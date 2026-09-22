"""Tests for the Energy Institute refinery capacity adapter.

NOTHING HERE TOUCHES THE NETWORK, and nothing here is a capture.

THE FIXTURE IS CONSTRUCTED IN CODE ON PURPOSE, AND THE REASON IS THE POINT OF
THE WHOLE MODULE. Every other adapter in this project is tested against a
verbatim slice of what the source served. This one cannot be: the Energy
Institute Statistical Review requires written permission for extensive
reproduction of its tables, and the capacity sheet is footnoted "Source:
Includes data from ICIS and S&P Global Energy" against a Review that calls
redistribution of S&P sourced data strictly prohibited. Committing a fixture
carrying real capacity numbers into a public repository would break the same
rule the adapter exists to respect, and it would do it in the test directory
where nobody would look for it.

So build_workbook below writes a workbook with the REAL SHAPE and INVENTED
NUMBERS. The shape is what these tests are about, and recon 03 section 2.2
measured every part of it against the real file:

    the sheet name, "Oil refinery - capacity"
    the units cell that opens the header row, "Thousand barrels daily*"
    years across the header from 1965
    the three columns headed 2025, of which only the first is a capacity, the
        others being a year on year growth rate and a share of world, marked by
        "Growth rate per annum" and "Share" in the row above
    country rows found by their printed labels, "United Kingdom" and not UK
    cells that are a mix of str and float, with floating point dust
    the sheet's own source footnote below the data

The numbers are a ramp. No test here asserts on a capacity value, because a
value that matched the real table would be the reproduction this file is
avoiding. What the tests assert is that the parser reads the right CELLS.
"""

from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

from crack.sources import ei
from crack.sources.base import SourceError

YEARS = list(range(1965, 2026))


def build_workbook(
    path: Path,
    *,
    years: list[int] | None = None,
    countries: list[str] | None = None,
    sheet: str = ei.SHEET,
    units_label: str = "Thousand barrels daily*",
    edition: str = "2026 Energy Institute Statistical Review of World Energy",
    growth_marker: str | None = "Growth rate per annum",
    repeat_label: str | None = None,
    missing_cell: tuple[str, int, object] | None = None,
) -> Path:
    """Write a workbook with the real shape and invented numbers. See the docstring."""
    import openpyxl

    years = list(YEARS if years is None else years)
    countries = list(COUNTRY_LABELS if countries is None else countries)

    book = openpyxl.Workbook()
    contents = book.active
    contents.title = ei.CONTENTS_SHEET
    if edition:
        contents.cell(row=2, column=1, value=edition)
        contents.cell(row=4, column=1, value="This workbook contains information presented in the")

    ws = book.create_sheet(sheet)
    ws.cell(row=1, column=1, value="Oil: Refining Capacity")
    # The real sheet, 1 indexed: the years end at column len(years) + 1, then
    # column len(years) + 2 carries the last year again over a growth rate,
    # len(years) + 3 a ten year CAGR headed "2015-25", and len(years) + 4 the
    # last year a third time over a share of world.
    if growth_marker:
        ws.cell(row=2, column=len(years) + 2, value=growth_marker)
        ws.cell(row=2, column=len(years) + 4, value="Share")
    ws.cell(row=3, column=1, value=units_label)
    for offset, year in enumerate(years):
        ws.cell(row=3, column=2 + offset, value=year)
    # The three columns headed by the same last year: the capacity is the one
    # inside the block, the other two are a growth rate and a share.
    if growth_marker:
        ws.cell(row=3, column=len(years) + 2, value=years[-1])
        ws.cell(row=3, column=len(years) + 3, value="%d-%d" % (years[-1] - 10, years[-1] % 100))
        ws.cell(row=3, column=len(years) + 4, value=years[-1])
        ws.cell(row=3, column=len(years) + 5, value=None)

    row = 5
    for index, label in enumerate(countries):
        ws.cell(row=row, column=1, value=label)
        for offset, year in enumerate(years):
            base = 100.0 * (index + 1) + (year - years[0])
            # A mix of str and float, and floating point dust, both real.
            value = str(int(base)) if year < 1980 else base + 1e-13
            ws.cell(row=row, column=2 + offset, value=value)
        if growth_marker:
            ws.cell(row=row, column=len(years) + 2, value=-0.078)
            ws.cell(row=row, column=len(years) + 4, value=0.012)
        row += 1
        if repeat_label == label:
            ws.cell(row=row, column=1, value=label)
            for offset, _ in enumerate(years):
                ws.cell(row=row, column=2 + offset, value=1.0)
            row += 1

    ws.cell(row=row + 1, column=1, value="Total Europe")
    ws.cell(row=row + 3, column=1, value="Source: Includes data from ICIS and S&P Global Energy")
    ws.cell(row=row + 4, column=1, value=" * Atmospheric distillation capacity at year end on a calendar-day basis.")
    ws.cell(row=row + 5, column=1, value=" n/a not available.")

    if missing_cell is not None:
        label, year, cell_value = missing_cell
        target_row = 5 + countries.index(label)
        ws.cell(row=target_row, column=2 + years.index(year), value=cell_value)

    book.save(path)
    return path


COUNTRY_LABELS = list(ei.COUNTRY_ROWS)


@pytest.fixture()
def workbook(tmp_path) -> Path:
    return build_workbook(tmp_path / "review.xlsx")


# --------------------------------------------------------------------------
# Parsing
# --------------------------------------------------------------------------

def test_the_five_countries_are_read_by_their_printed_labels(workbook):
    frame = ei.parse_capacity(workbook)
    assert list(frame.columns) == [
        "date",
        "be_capacity_kb_d",
        "de_capacity_kb_d",
        "fr_capacity_kb_d",
        "nl_capacity_kb_d",
        "gb_capacity_kb_d",
        "nwe5_capacity_kb_d",
    ]
    assert len(frame) == len(YEARS)


def test_the_label_is_united_kingdom_and_the_column_is_gb():
    """EI spells it out, JODI codes it GB, SPEC.md calls it UK in prose."""
    assert ei.COUNTRY_ROWS["United Kingdom"] == "GB"
    assert "UK" not in ei.COUNTRY_ROWS
    assert ei.column_name("GB") == "gb_capacity_kb_d"


def test_the_date_is_the_year_end_because_the_figure_is_a_year_end_stock(workbook):
    frame = ei.parse_capacity(workbook)
    assert frame["date"].iloc[0] == pd.Timestamp("1965-12-31")
    assert frame["date"].iloc[-1] == pd.Timestamp("2025-12-31")


def test_the_growth_rate_column_is_not_read_as_a_capacity(workbook):
    """The header prints the last year three times and only the first is capacity.

    A parser that took every integer header would read a growth rate of minus
    0.078 as a capacity of minus 0.078 thousand barrels a day, which is both
    wrong and, being small and negative, easy to miss on a chart of 6,000.
    """
    frame = ei.parse_capacity(workbook)
    last = frame.iloc[-1]
    for code in ei.COUNTRIES:
        assert last[ei.column_name(code)] > 1.0
    # One row per year and not one per year plus three.
    assert frame["date"].dt.year.tolist() == YEARS


def test_a_year_headed_twice_inside_the_block_is_a_stop(tmp_path):
    path = build_workbook(tmp_path / "dup.xlsx", years=YEARS + [2025])
    with pytest.raises(SourceError, match="headed twice"):
        ei.parse_capacity(path)


def test_string_cells_are_coerced_and_dust_is_left_alone(workbook):
    """Many 1960s and 1970s cells arrive as strings, and later ones carry dust."""
    frame = ei.parse_capacity(workbook)
    early = frame.loc[frame["date"].dt.year == 1970, "be_capacity_kb_d"].iloc[0]
    assert float(early) == pytest.approx(105.0)
    late = frame.loc[frame["date"].dt.year == 2000, "be_capacity_kb_d"].iloc[0]
    assert isinstance(late, float)


def test_a_footnote_token_becomes_nan_and_never_a_zero(tmp_path):
    """"^ Less than 0.5" is a suppressed figure, not a capacity of nothing."""
    path = build_workbook(
        tmp_path / "na.xlsx", missing_cell=("Belgium", 1975, "n/a")
    )
    frame = ei.parse_capacity(path)
    value = frame.loc[frame["date"].dt.year == 1975, "be_capacity_kb_d"].iloc[0]
    assert np.isnan(value)


def test_the_five_country_total_is_nan_unless_all_five_printed_a_number(tmp_path):
    path = build_workbook(tmp_path / "hole.xlsx", missing_cell=("France", 1990, "n/a"))
    frame = ei.parse_capacity(path)
    row = frame.loc[frame["date"].dt.year == 1990].iloc[0]
    assert np.isnan(row["nwe5_capacity_kb_d"])
    assert not np.isnan(row["be_capacity_kb_d"])


def test_an_unreadable_cell_raises_rather_than_becoming_a_number(tmp_path):
    path = build_workbook(tmp_path / "odd.xlsx", missing_cell=("Germany", 1995, "circa 2000"))
    with pytest.raises(SourceError, match="cannot read"):
        ei.parse_capacity(path)


def test_a_missing_country_row_is_a_stop(tmp_path):
    path = build_workbook(
        tmp_path / "gone.xlsx",
        countries=[label for label in COUNTRY_LABELS if label != "Netherlands"],
    )
    with pytest.raises(SourceError, match="no row labelled 'Netherlands'"):
        ei.parse_capacity(path)


def test_a_country_printed_twice_is_a_stop(tmp_path):
    path = build_workbook(tmp_path / "twice.xlsx", repeat_label="France")
    with pytest.raises(SourceError, match="twice"):
        ei.parse_capacity(path)


def test_a_renamed_sheet_is_a_stop(tmp_path):
    path = build_workbook(tmp_path / "renamed.xlsx", sheet="Oil refining capacity")
    with pytest.raises(SourceError, match="no sheet named"):
        ei.parse_capacity(path)


def test_a_missing_units_row_is_a_stop(tmp_path):
    path = build_workbook(tmp_path / "nounits.xlsx", units_label="Kilotonnes")
    with pytest.raises(SourceError, match="no header row"):
        ei.parse_capacity(path)


def test_a_history_that_starts_later_than_1965_is_a_stop(tmp_path):
    path = build_workbook(tmp_path / "short.xlsx", years=list(range(1990, 2026)))
    with pytest.raises(SourceError, match="now starts at 1990"):
        ei.parse_capacity(path)


# --------------------------------------------------------------------------
# The edition, read out of the file
# --------------------------------------------------------------------------

def test_the_edition_comes_from_the_contents_sheet(workbook):
    assert ei.read_edition(workbook) == (
        "2026 Energy Institute Statistical Review of World Energy"
    )


def test_a_workbook_with_no_edition_line_is_a_stop(tmp_path):
    path = build_workbook(tmp_path / "anon.xlsx", edition="")
    with pytest.raises(SourceError, match="does not print an edition line"):
        ei.read_edition(path)


# --------------------------------------------------------------------------
# Monthly capacity and utilisation, prepared for whichever option is taken
# --------------------------------------------------------------------------

def capacity_frame() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "date": [pd.Timestamp("2024-12-31"), pd.Timestamp("2025-12-31")],
            "nwe5_capacity_kb_d": [6000.0, 6120.0],
        }
    )


def test_step_capacity_holds_the_year_end_across_its_own_year():
    monthly = ei.capacity_monthly(capacity_frame(), how="step")
    # December 2024 to December 2025 inclusive.
    assert len(monthly) == 13
    assert monthly["nwe5_capacity_kb_d"].iloc[0] == 6000.0
    assert monthly.loc[monthly["date"] == pd.Timestamp("2024-12-01"), "nwe5_capacity_kb_d"].iloc[0] == 6000.0
    assert monthly.loc[monthly["date"] == pd.Timestamp("2025-01-01"), "nwe5_capacity_kb_d"].iloc[0] == 6120.0


def test_linear_capacity_walks_between_the_year_ends():
    monthly = ei.capacity_monthly(capacity_frame(), how="linear")
    first = monthly["nwe5_capacity_kb_d"].iloc[0]
    last = monthly["nwe5_capacity_kb_d"].iloc[-1]
    middle = monthly.loc[
        monthly["date"] == pd.Timestamp("2025-06-01"), "nwe5_capacity_kb_d"
    ].iloc[0]
    assert first == 6000.0
    assert last == 6120.0
    assert 6000.0 < middle < 6120.0


def test_capacity_is_never_extended_past_the_last_published_year(tmp_path):
    """There is no 2026 capacity anywhere, and holding 2025 forward is an assumption.

    It is a defensible one and it is not a published figure, so it belongs in the
    analysis where it can be labelled, not in a function that looks like it read
    something.
    """
    monthly = ei.capacity_monthly(capacity_frame())
    assert monthly["date"].max() == pd.Timestamp("2025-12-01")


def test_utilisation_is_nan_where_no_capacity_was_published():
    intake = pd.DataFrame(
        {
            "date": pd.date_range("2025-11-01", periods=4, freq="MS"),
            "nwe5_refinobs_crudeoil_kbd": [5000.0, 5100.0, 5200.0, 5300.0],
        }
    )
    out = ei.utilisation(intake, capacity_frame())
    assert out["utilisation"].iloc[0] == pytest.approx(5000.0 / 6120.0)
    assert out["capacity_is_published"].tolist() == [True, True, False, False]
    assert np.isnan(out["utilisation"].iloc[2])


def test_utilisation_refuses_a_column_that_is_not_there():
    with pytest.raises(SourceError, match="no column"):
        ei.utilisation(
            pd.DataFrame({"date": [pd.Timestamp("2025-01-01")], "other": [1.0]}),
            capacity_frame(),
        )


def test_capacity_monthly_refuses_an_unknown_interpolation():
    with pytest.raises(ValueError, match="step"):
        ei.capacity_monthly(capacity_frame(), how="spline")


# --------------------------------------------------------------------------
# The adapter
# --------------------------------------------------------------------------

def adopted(sandbox, workbook: Path) -> ei.EiRefineryCapacity:
    adapter = ei.EiRefineryCapacity(root=sandbox / "private" / "ei")
    adapter.adopt_workbook(
        workbook,
        obtained_at="2026-09-11T00:00:00Z",
        note="constructed by tests/test_sources_ei.py, see its docstring",
    )
    return adapter


def test_the_cache_lands_in_private_and_never_in_cache(sandbox, workbook):
    """SPEC.md non negotiable 6. The flag is what keeps the bytes out of the commit."""
    entry = adopted(sandbox, workbook).run()
    assert entry["committable"] is False
    assert entry["file"] == "data/private/ei_refinery_capacity_annual.csv"
    assert (sandbox / "private" / "ei_refinery_capacity_annual.csv").exists()
    assert not (sandbox / "cache" / "ei_refinery_capacity_annual.csv").exists()


def test_the_manifest_states_the_prohibition_in_words(sandbox, workbook):
    entry = adopted(sandbox, workbook).run()
    assert "strictly prohibited" in entry["prohibition"]
    assert "statisticalreview@energyinst.org" in entry["prohibition"]
    assert "S&P Global Energy" in entry["sheet_source_footnote"]
    assert entry["licence_note"]


def test_the_manifest_says_which_year_the_denominator_is_from(sandbox, workbook):
    """No 2026 capacity exists anywhere, so nothing may imply the figure is current."""
    entry = adopted(sandbox, workbook).run()
    assert entry["latest_year"] == 2025
    assert entry["last_date"] == "2025-12-31"
    assert "no 2026 capacity" in entry["currency_note"].lower()
    assert "THE LATEST YEAR IS 2025" in entry["note"]


def test_the_manifest_leaves_the_gate_1_decision_open(sandbox, workbook):
    entry = adopted(sandbox, workbook).run()
    decision = entry["gate_1_decision"]
    assert decision["chosen"] is None
    assert len(decision["options"]) == 3


def test_a_run_that_fetched_nothing_does_not_print_a_fresh_fetch_time(sandbox, workbook):
    """fetched_at means when the bytes were obtained, not when this program ran."""
    entry = adopted(sandbox, workbook).run()
    assert entry["fetched_at"] == "2026-09-11T00:00:00Z"
    assert entry["checked_at"] > entry["fetched_at"]
    assert entry["workbook"]["origin"] == "adopted"
    assert entry["workbook"]["sha256"]
    assert entry["workbook_fetched_this_run"] is False


def test_the_manifest_records_that_the_live_fetch_is_blocked(sandbox, workbook):
    entry = adopted(sandbox, workbook).run()
    assert entry["live_fetch"].startswith("BLOCKED")
    assert "Cf-Mitigated" in entry["live_fetch"]


def test_a_blocked_fetch_says_what_to_do_and_does_not_try_to_get_around_it(
    sandbox, monkeypatch, tmp_path
):
    def refuse(*args, **kwargs):
        raise SourceError("GET %s failed with HTTP 403, not retryable" % args[0])

    monkeypatch.setattr(ei, "http_get", refuse)
    adapter = ei.EiRefineryCapacity(root=tmp_path / "empty")
    with pytest.raises(SourceError) as caught:
        adapter.load_workbook_bytes()
    message = str(caught.value)
    assert "Cf-Mitigated: challenge" in message
    assert "Download EI-Stats-Review-ALL-data.xlsx by hand" in message


def test_an_interstitial_saved_under_an_xlsx_name_is_refused(sandbox, tmp_path):
    decoy = tmp_path / "EI-Stats-Review-ALL-data.xlsx"
    decoy.write_bytes(b"<!DOCTYPE html><html><title>Just a moment...</title>")
    adapter = ei.EiRefineryCapacity(root=sandbox / "private" / "ei")
    with pytest.raises(SourceError, match="not a zip container"):
        adapter.adopt_workbook(decoy, obtained_at="2026-09-13T00:00:00Z", note="decoy")


def test_a_failed_parse_keeps_the_previous_cache(sandbox, tmp_path, workbook):
    good = adopted(sandbox, workbook)
    good.run()
    before = (sandbox / "private" / "ei_refinery_capacity_annual.csv").read_bytes()

    broken = build_workbook(tmp_path / "broken.xlsx", years=list(range(1990, 2026)))
    adapter = ei.EiRefineryCapacity(root=sandbox / "private" / "ei")
    adapter.adopt_workbook(broken, obtained_at="2026-09-13T00:00:00Z", note="broken")
    with pytest.raises(SourceError):
        adapter.run()

    assert (sandbox / "private" / "ei_refinery_capacity_annual.csv").read_bytes() == before
    manifest = json.loads((sandbox / "manifest.json").read_text(encoding="utf-8"))
    entry = next(e for e in manifest["series"] if e["series"] == ei.SERIES)
    assert entry["status"] == "failed"


def test_the_utilisation_path_works_end_to_end_on_the_constructed_workbook(
    sandbox, workbook
):
    """Option 2 needs no new parsing. This proves the two frames join.

    It does not choose option 2. See the module docstring on why publishing the
    ratio does not by itself withhold the table.
    """
    adapter = adopted(sandbox, workbook)
    capacity = adapter.fetch()
    intake = pd.DataFrame(
        {
            "date": pd.date_range("2025-01-01", periods=12, freq="MS"),
            "nwe5_refinobs_crudeoil_kbd": np.linspace(4000.0, 4400.0, 12),
        }
    )
    out = ei.utilisation(intake, capacity)
    assert out["capacity_is_published"].all()
    assert (out["utilisation"] > 0).all()


# --------------------------------------------------------------------------
# The one derived series that IS committed
# --------------------------------------------------------------------------
#
# Gate 5. A GitHub runner has a fresh clone and nothing from data/private, so
# before this series existed the analysis could not run there at all: 98 tests
# and python scripts/export.py --check died on FileNotFoundError. The five
# country total is now committed and every per country row is still withheld.
# These tests hold both halves of that down. See crack.sources.ei's docstring,
# "WHAT GATE 5 ADDED".


def derived(sandbox, workbook) -> ei.Nwe5RefineryCapacity:
    """The private table written, then the derived total built from it."""
    adopted(sandbox, workbook).run()
    return ei.Nwe5RefineryCapacity(root=sandbox / "private" / "ei")


def test_the_derived_total_lands_in_cache_and_is_committable(sandbox, workbook):
    entry = derived(sandbox, workbook).run()
    assert entry["committable"] is True
    assert entry["file"] == "data/cache/nwe5_refinery_capacity_annual.csv"
    assert (sandbox / "cache" / "nwe5_refinery_capacity_annual.csv").exists()
    assert not (sandbox / "private" / "nwe5_refinery_capacity_annual.csv").exists()
    # The private table is still private, and still the only thing parsed.
    assert (sandbox / "private" / "ei_refinery_capacity_annual.csv").exists()


def test_the_derived_cache_carries_the_total_and_no_country_row(sandbox, workbook):
    """THE LICENCE PROMISE, in the file itself. One column beside the date."""
    derived(sandbox, workbook).run()
    text = (sandbox / "cache" / "nwe5_refinery_capacity_annual.csv").read_text(
        encoding="utf-8"
    )
    header = text.splitlines()[0].split(",")
    assert header == ["date", ei.column_name(ei.AGGREGATE)]
    for code in ei.COUNTRIES:
        assert ei.column_name(code) not in text


def test_the_derived_total_is_the_sum_of_the_five_private_rows(sandbox, workbook):
    """Derived, not re-parsed: the same arithmetic, checked against the table."""
    private = adopted(sandbox, workbook).fetch()
    total = ei.Nwe5RefineryCapacity(
        source_frame=private, root=sandbox / "private" / "ei"
    ).fetch()
    assert list(total.columns) == ["date", ei.column_name(ei.AGGREGATE)]
    assert len(total) == len(private)
    pd.testing.assert_series_equal(
        total[ei.column_name(ei.AGGREGATE)],
        private[ei.column_name(ei.AGGREGATE)],
        check_names=False,
    )


def test_the_manifest_entry_names_neither_a_country_column_nor_the_country_values(
    sandbox, workbook
):
    """The manifest is rendered in full on the provenance panel, so it is a
    committed file like any other and tools/validate-data.mjs check 16 scans it."""
    entry = derived(sandbox, workbook).run()
    blob = json.dumps(entry)
    for code in ei.COUNTRIES:
        assert ei.column_name(code) not in blob
    assert entry["published_columns"] == [ei.column_name(ei.AGGREGATE)]
    assert entry["derived_from"] == ei.SERIES
    assert entry["method"] == "derived"
    assert "THE FIVE COUNTRY TOTAL ONLY" in entry["note"]


def test_it_refuses_to_guess_when_the_private_table_is_not_on_the_machine(sandbox):
    """CI has no data/private. It reads the committed cache and never runs this."""
    with pytest.raises(SourceError) as raised:
        ei.Nwe5RefineryCapacity(root=sandbox / "private" / "ei").run()
    message = str(raised.value)
    assert "data/private/ei_refinery_capacity_annual.csv" in message
    assert "not on this machine" in message
    assert not (sandbox / "cache" / "nwe5_refinery_capacity_annual.csv").exists()
