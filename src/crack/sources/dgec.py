"""The two DGEC monthly historical workbooks, the official spine of this study.

    series      dgec_brent_monthly        dgec_mbr_monthly
    cache       data/cache/<series>.csv
    columns     date, brent_usd_bbl,      date, mbr_usd_bbl, mbr_eur_t
                brent_usd_t
    unit        USD per barrel, plus a    USD per barrel and EUR per tonne
                USD per tonne column
                derived at 7.5 bbl/t
    source      DGEC, Direction generale de l'energie et du climat, ministere de
                la Transition ecologique, assessments credited to Reuters
    page        https://www.ecologie.gouv.fr/politiques-publiques/prix-produits-petroliers

These two files are why this project has a margin layer at all. Recon 02 section
0 put it bluntly: the ministry publishes no historical file for the Rotterdam
product quotations, keeps exactly one weekly note online at a time and deletes
the previous one, and the Internet Archive holds nine notes for the whole 2024 to
2026 period. What does exist, machine readable and complete, is these two
workbooks: Brent date monthly since January 2015, and the official gross refining
margin on Brent, the marge brute de raffinage sur Brent or MBR, monthly since
January 2015. Every one of the eight MBR anchors in SPEC.md section 5.5
reproduces from the second of them exactly, which is what makes
tests/test_dgec_anchor.py a gate rather than a formality.

The weekly note itself, its printed $/t quotation table and its chart, is a
different adapter and a different batch. Nothing here parses a PDF.

THE URL CANNOT BE HARDCODED, AND THE LABEL CANNOT BE TRUSTED EITHER
-------------------------------------------------------------------
SPEC.md section 5.6 lists the ministry page as
http://www.developpement-durable.gouv.fr/prix-des-produits-petroliers-1 . Recon
02 section 1.1 measured it: HTTP 404, after a redirect to ecologie.gouv.fr. The
live page is LANDING_PAGE below.

On that page the two workbooks are linked by Drupal, which appends a duplicate
filename counter, _0, _1 and so on, at upload time. Recon 02 section 1.2 found
the href and the visible label disagreeing about that suffix IN BOTH DIRECTIONS
on the same page, on the same day:

    Brent   href ends _0.xlsx                 label ends .xlsx
    MBR     href ends (moyennes mensuelles)   label ends _1.xlsx
            .xlsx, no suffix

and the _1 spelling of the MBR URL, the one the label claims, returns HTTP 404.
Wayback shows the same two files captured under three different spellings over
two years. So:

    the href is authoritative, the label is not, and neither is stable.

discover_workbook_url() therefore reads the landing page, decodes each href,
strips the extension and any _N counter, and matches what is left against the
document's title. It RAISES when it finds no match or more than one. It never
falls back to a remembered URL, for the reason worldbank.py gives for the same
decision: a remembered URL that still resolves is the worst available outcome,
an old edition fetched successfully and stamped with a vintage nobody read.

COLUMN A IS NOT ALWAYS A DATE
------------------------------
Both workbooks carry an ANNUAL AVERAGES block below the monthly one, keyed by a
year written as TEXT: '2015', '2016' and so on. Coerced with pandas defaults,
'2015' becomes 2015-01-01 and collides with the real January 2015 row. Recon 02
sections 1.3 and 1.4 documented the block; SPEC.md section 13 says a layout like
this must break loudly rather than quietly return the wrong thing.

The parser handles it by bounding the block rather than by filtering rows:
_monthly_block() finds the header row whose column A reads "Moyennes mensuelles",
finds the row whose column A reads "MOYENNES ANNUELLES", takes what is strictly
between them, and then REQUIRES every non empty column A cell inside that span to
be a real date. A text year inside the monthly block is an error, not a row to
skip. That way, if the ministry ever drops the annual heading or moves the annual
block above the monthly one, the parser stops with the offending cell named
instead of silently absorbing eleven annual averages into a monthly series.

THE TWO BRENT CONVERSION FACTORS, AND WHY THIS MODULE USES 7.5
---------------------------------------------------------------
DGEC uses two different barrels per tonne figures for Brent, in two different
places. Both are in crack.config with their citations:

    DGEC_BBL_PER_T_BRENT_NOTE   = 7.5    the factor the weekly note uses to
                                         print Brent date in its $/t quotation
                                         table. Not stated anywhere. Recon 02
                                         section 4.3 recovered it by dividing
                                         the note's printed $/t by the $/bbl
                                         DGEC publishes for the same month in
                                         the workbook this module reads, on five
                                         month pairs, and got 7.5 exactly each
                                         time.
    DGEC_BBL_PER_T_BRENT_MARGIN = 7.55   stated in the MBR methodology note,
                                         page 4, for the CAF Brent price INSIDE
                                         the margin calculation.

This module uses 7.5 and only 7.5, because the one thing brent_usd_t is for is
to sit next to the note's printed $/t quotations, where the products are quoted,
and to reproduce the SPEC.md section 5.5 anchor of 628 $/t for July 2026. Using
7.55 there would be wrong by 0.67 percent, about 4 $/t at the 2026 level, and it
would be silent. 7.55 belongs to a replication of the MBR and to nothing else,
and recon 02 section 5.6 found that replication cannot in fact be completed from
published data because five of the ten product quotations, the PEG Nord gas
price and the Aframax freight are never published.

brent_usd_t is DERIVED. It is one multiplication, declared in the manifest under
derived_columns with its factor and its citation, and the site must label it that
way. The note prints the same quantity rounded to the whole dollar.

PROVISIONAL, AND THE HONEST ANSWER ABOUT VINTAGE
-------------------------------------------------
SPEC.md non negotiable 5 says the current month is provisional and gets revised,
and recon 02 section 3.5 measured how much: the note whose content date is 9
August 2024 printed August 2024 at 5.63 $/b provisional, the note of 6 September
2024 printed it at 3.63 final, and the workbook now says 3.6331. A 2 $/b swing.
Never blend a provisional month into a final series.

What that means HERE is narrower than it sounds, and saying so precisely is the
point. These workbooks carry complete months only: read on 2026-09-11 the Brent
file ended 2026-08 and the MBR file ended 2026-07, while the note of 4 September
2026 already printed August 2026 at 38.05. The first live fetch by this adapter,
on 2026-09-12, got an MBR workbook stamped Last-Modified Fri, 11 Sep 2026
14:52:54 GMT that had caught up to 2026-08 at 38.0505047, agreeing with the
note's printed figure. So the provisional figure lives in the note, not in these
files, and the lag between the two is a scheduling artefact rather than a rule.
provisional_from is therefore set mechanically, by
provisional_month(): the last month in the file is flagged provisional only when
it is the month the fetch is happening in, which is the only circumstance under
which DGEC would be publishing an incomplete month here. Otherwise it is None and
the manifest note says why, rather than a blank field a reader has to interpret.

THE MBR FILE LAGS THE NOTE BY AT LEAST A MONTH. Recon 02 section 1.5. An adapter
that reads only the xlsx is a month stale on the headline margin, and the
manifest carries note_leads_file saying so. Splicing the note's last complete
month onto the workbook is recon 02 open question 5 and is NOT done here: it
would mix two vintages of the same series under one label, which is exactly what
non negotiable 5 forbids. When the note adapter exists, the splice is a decision
for the owner and it gets its own flagged column.

The workbooks carry no "updated on" cell, so there is no vintage printed inside
them the way the pink sheet prints one. vintage is assembled from the two facts
that do exist: the last month present, and the HTTP Last-Modified header if the
server sends one.

WHAT IS NOT IN THESE FILES
---------------------------
SPEC.md section 5.1 says the MBR is published "in $/bbl, EUR/t and c EUR/l". The
workbook carries two of the three. There is no centimes per litre column; recon
02 section 1.4 checked. That unit exists only in the weekly note, whose MBR table
prints en $/b, en EUR/t and en c EUR/l. This module does not compute one, because
computing it would need a litres per tonne density for the product basket that
DGEC does not publish, and the result would be this study's number wearing the
ministry's label. The gap travels in the manifest as units_not_in_file.

THE METHOD CHANGE SPEC.md SECTION 4.3 ASKS TO MARK IS NOT INSIDE THIS WINDOW
-----------------------------------------------------------------------------
The methodology note says the MBR has been published since 1998, that the method
was revised with IFPEN in 2014, that the detailed method in its annex is the one
used from 1 January 2016, and that 2014 and 2015 were recomputed retroactively on
the new basis for statistical continuity. The published file starts in January
2015, which is INSIDE the recomputed window, so the whole series available to
this project is on one method and no chart it can draw spans the break. That is
carried in the manifest as method_change rather than left for the Method view to
assert, and recon 02 section 5.1 has the verbatim quotations.

LICENCE
--------
Licence Ouverte 2.0, Etalab. It expressly permits extraction, transformation,
redistribution and publication, including commercially, on condition of
attribution to the concedant and a statement of the date of last update. Both
tables are credited "Source : DGEC-REUTERS" and the underlying assessments are a
commercial vendor's, which recon 02 section 6.4 raises as a real doubt rather
than waving through. The mitigation is in crack.config: republish the parsed
values with attribution, never the PDFs.
"""

from __future__ import annotations

import datetime as _dt
import io
import re
import time
import unicodedata
from html.parser import HTMLParser
from typing import Iterable, Sequence
from urllib.parse import unquote, urljoin

import pandas as pd

from ..config import (
    BOUNDS_BRENT_USD_BBL,
    BOUNDS_MARGIN_USD_BBL,
    BOUNDS_PRODUCT_USD_T,
    DGEC_BBL_PER_T_BRENT_MARGIN,
    DGEC_BBL_PER_T_BRENT_NOTE,
)
from ..config import source as registered_source
from .base import Adapter, SourceError, http_get

__all__ = [
    "SERIES_BRENT",
    "SERIES_MBR",
    "COLUMN_BRENT_USD_BBL",
    "COLUMN_BRENT_USD_T",
    "COLUMN_MBR_USD_BBL",
    "COLUMN_MBR_EUR_T",
    "LANDING_PAGE",
    "BRENT_DOCUMENT_TITLE",
    "MBR_DOCUMENT_TITLE",
    "BRENT_SHEET",
    "MBR_SHEET",
    "BRENT_TITLE_CELL",
    "MBR_TITLE_CELL",
    "MONTHLY_BLOCK_LABEL",
    "ANNUAL_BLOCK_LABEL",
    "UNIT_USD_BBL",
    "UNIT_EUR_T",
    "SERIES_START",
    "BOUNDS_MBR_EUR_T",
    "METHOD_CHANGE",
    "NOTE_LEADS_FILE",
    "UNITS_NOT_IN_FILE",
    "normalise",
    "discover_workbook_url",
    "parse_brent_workbook",
    "parse_mbr_workbook",
    "provisional_month",
    "clear_landing_cache",
    "fetch_landing_page",
    "fetch_brent",
    "fetch_mbr",
    "DgecBrentMonthly",
    "DgecMbrMonthly",
    "main",
]


# --------------------------------------------------------------------------
# What to look for, all of it a label and none of it a position
# --------------------------------------------------------------------------

#: Manifest and cache names. Both are keys of crack.config.SOURCES, and
#: Adapter._check_declarations cross checks this module against that registry on
#: frequency, method and committable before anything is fetched.
SERIES_BRENT = "dgec_brent_monthly"
SERIES_MBR = "dgec_mbr_monthly"

#: Value columns. Units are in the names because every consumer of this cache
#: reads them straight into a formula, and a renamed column would hide a unit
#: change. brent_usd_t is DERIVED, see the module docstring.
COLUMN_BRENT_USD_BBL = "brent_usd_bbl"
COLUMN_BRENT_USD_T = "brent_usd_t"
COLUMN_MBR_USD_BBL = "mbr_usd_bbl"
COLUMN_MBR_EUR_T = "mbr_eur_t"

#: The page the two workbook links are discovered on. This one is stable. The
#: file URLs are not, see the module docstring.
LANDING_PAGE = (
    "https://www.ecologie.gouv.fr/politiques-publiques/prix-produits-petroliers"
)

#: The document titles, as the href spells them once it is percent decoded and
#: once the extension and any Drupal _N counter are stripped. Matched after
#: normalise(), so accents and case do not matter and the file can be named with
#: or without them.
BRENT_DOCUMENT_TITLE = "Historique du cours du Brent depuis 2015 (moyennes mensuelles)"
MBR_DOCUMENT_TITLE = (
    "Historique de la marge brute de raffinage sur Brent depuis 2015 "
    "(moyennes mensuelles)"
)

#: Sheet names, verbatim. Each workbook holds exactly one sheet.
BRENT_SHEET = "Cours du Brent"
MBR_SHEET = "Marge brute raffinage sur Brent"

#: Cell B1 of each sheet, normalised. Checked so that a workbook served under the
#: right filename but holding some other table is refused rather than parsed.
BRENT_TITLE_CELL = "cours moyens mensuels du brent date en $/baril"
MBR_TITLE_CELL = "marge brute de raffinage sur brent"

#: Column A of the header row of the monthly block, and of the heading that ends
#: it. Both normalised. See _monthly_block().
MONTHLY_BLOCK_LABEL = "moyennes mensuelles"
ANNUAL_BLOCK_LABEL = "moyennes annuelles"

#: The unit cells on the header row, normalised. normalise() rewrites the euro
#: sign as "eur" so this module stays ASCII. The unit is checked for every column
#: read, because a series that changed unit would be wrong by a factor of about
#: seven and nothing downstream would notice.
UNIT_USD_BBL = "en $/b"
UNIT_EUR_T = "en eur/t"

#: The first month both files carry. A workbook starting later has been re cut
#: and is worth stopping on rather than silently shortening the history, which is
#: the rule worldbank.py already applies to the pink sheet.
SERIES_START = pd.Timestamp("2015-01-01")

#: Row counts when recon 02 read the files on 2026-09-11: 140 monthly rows of
#: Brent, 2015-01 to 2026-08, and 139 of MBR, 2015-01 to 2026-07. The floors sit
#: under those with a little room, and high enough that a truncated workbook
#: cannot replace a good cache. The Adapter additionally refuses any frame with
#: fewer rows than the cache already on disk, so these only have to catch the
#: first fetch.
MIN_ROWS_BRENT = 130
MIN_ROWS_MBR = 130

#: Bounds for the MBR in EUR per tonne, the one value column in this module with
#: no bound already named in crack.config.
#:
#: It is DERIVED FROM ONE, not invented: the $/bbl band, crack.config
#: BOUNDS_MARGIN_USD_BBL of (-50, 150), converted at DGEC's own margin factor
#: DGEC_BBL_PER_T_BRENT_MARGIN of 7.55 bbl/t and read at a euro of one dollar,
#: giving (-377.5, 1132.5), rounded outwards to whole hundreds. Like every other
#: bound in this project it is a units and sign check and not a view on the
#: market: the observed range of this column over 2015-01 to 2026-08 is -1.79 to
#: 246.11 EUR/t, so the band is roughly five times the observed range and only a
#: file that arrived in a different unit or with a decimal point in the wrong
#: place can trip it.
BOUNDS_MBR_EUR_T = (-400.0, 1200.0)

#: SPEC.md section 4.3 asks for the MBR method change date and says to mark it as
#: a structural break on every chart that spans it. It is carried here, in the
#: manifest, because the honest answer is that no chart this project can draw
#: spans it. Quotations are verbatim from the methodology note, recon 02 section
#: 5.1.
METHOD_CHANGE = {
    "published_since": "1998",
    "revised": "2014, with IFPEN",
    "detailed_method_in_force_from": "2016-01-01",
    "recomputed_backwards_over": "2014 to 2015",
    "break_inside_the_published_window": False,
    "note": (
        "The methodology note says 'La DGEC calcule et diffuse une marge brute "
        "de raffinage sur Brent depuis 1998', that the method was revised with "
        "IFPEN in 2014, that 'Le detail du mode de calcul de la marge brute "
        "utilise a partir du 1er janvier 2016 figure en annexe', and that 'Afin "
        "d'assurer la continuite statistique, la marge de raffinage a ete "
        "retroactivement recalculee sur la base de la nouvelle methodologie pour "
        "la periode 2014-2015'. The published file starts 2015-01, inside that "
        "recomputed window, so the entire series available to this project is on "
        "one method and contains no break to mark. SPEC.md section 4.3 expects a "
        "break; there is none to draw, and the Method view says so rather than "
        "drawing a line where nothing happened. The pre 2014 method is not "
        "comparable and the pre 2015 data is not published."
    ),
    "source": (
        "Mode de calcul de la marge brute de raffinage sur brent, edition of "
        "1 August 2019, https://www.ecologie.gouv.fr/sites/default/files/"
        "documents/Mode%20de%20calcul%20de%20la%20marge%20brute%20de%20"
        "raffinage%20sur%20brent.pdf"
    ),
}

#: Recon 02 section 1.5, carried into the manifest so the provenance panel can
#: say it rather than leaving a reader to notice the last date is old.
NOTE_LEADS_FILE = (
    "The weekly note can lead this workbook by a month, and how long it does is "
    "not fixed. Two measurements, a day apart, both recorded because the second "
    "does not cancel the first. On 2026-09-11 recon 02 section 1.5 read an MBR "
    "workbook ending 2026-07 while the note of 4 September 2026 already printed "
    "August 2026 at 38.05 $/b, final, in the column carrying no '(donnees "
    "provisoires)' marker. On 2026-09-12 the first live fetch by this adapter got "
    "a workbook stamped Last-Modified Fri, 11 Sep 2026 14:52:54 GMT, ending "
    "2026-08 at 38.0505047 $/b, which agrees with the note's printed 38,05 to the "
    "two decimals the note prints. So the ministry closed the gap within a day, "
    "and the lag is a scheduling artefact rather than a rule: expect the workbook "
    "to be behind the note for part of each month. The splice of the note's last "
    "complete month onto the workbook is recon 02 open question 5 and is NOT done "
    "here: it would mix two vintages of one series under one label. When it is "
    "done it gets its own flagged column."
)

#: SPEC.md section 5.1 expects three units and the file carries two.
UNITS_NOT_IN_FILE = (
    "SPEC.md section 5.1 says the MBR is published in $/bbl, EUR/t and c EUR/l. "
    "This workbook carries $/b and EUR/t only; there is no centimes per litre "
    "column, recon 02 section 1.4. That unit exists solely in the weekly note, "
    "whose MBR table prints all three. None is computed here: it would need a "
    "litres per tonne density for DGEC's product basket that DGEC does not "
    "publish, and the result would be this study's number wearing the ministry's "
    "label."
)

#: How long a landing page fetched by one adapter may be reused by the other, in
#: seconds. Both series live on one page and SPEC.md section 5.4 asks for one
#: request per file and a polite delay between requests, so refetching the same
#: HTML twice in the same run would be rude for no gain. Five minutes is long
#: enough for a refresh run and far too short to hide an edition change.
LANDING_CACHE_SECONDS = 300.0


_registered_brent = registered_source(SERIES_BRENT)
_registered_mbr = registered_source(SERIES_MBR)
for _reg in (_registered_brent, _registered_mbr):
    if _reg.page_url != LANDING_PAGE:
        raise RuntimeError(
            "crack.sources.dgec discovers both workbooks from %s but "
            "crack.config.SOURCES registers %s for %s"
            % (LANDING_PAGE, _reg.page_url, _reg.series)
        )
    if _reg.machine_url is not None:
        raise RuntimeError(
            "crack.config.SOURCES gives %s a machine_url of %r. The Drupal _N "
            "counter in that path is unstable and the URL the visible label "
            "claims returns HTTP 404, recon 02 section 1.2, so the link is "
            "discovered from the landing page and there is no URL to register"
            % (_reg.series, _reg.machine_url)
        )


# --------------------------------------------------------------------------
# Text handling
# --------------------------------------------------------------------------

#: Whitespace this source uses that is not a plain space. The workbooks and the
#: notes both use a narrow no break space as a thousands separator and a no break
#: space before a colon, and either one inside a label would defeat an exact
#: comparison.
_SPACE_CODEPOINTS = (
    0x00A0,  # no break space
    0x2007,  # figure space
    0x2009,  # thin space
    0x202F,  # narrow no break space, the notes' thousands separator
    0x2002,  # en space
    0x2003,  # em space
    0x0009,  # tab
    0x000A,  # newline
    0x000D,  # carriage return
)
_SPACES = dict.fromkeys(_SPACE_CODEPOINTS, " ")

#: Currency signs rewritten rather than dropped, so this module stays ASCII while
#: still telling a euro column from a dollar one. A normalisation that discarded
#: the sign would let "en EUR/t" match a unit cell reading "en $/t".
_CURRENCY = {0x20AC: "eur", 0x00A3: "gbp"}


def normalise(text) -> str:
    """Fold a cell or a filename to a comparable ASCII form.

    Accents are stripped, the euro sign becomes "eur", the several kinds of non
    breaking space become ordinary spaces, runs of whitespace collapse, and the
    result is casefolded and trimmed.

    This exists so that every label in this module can be written in plain ASCII
    while still matching a French workbook exactly. It is deliberately NOT a
    fuzzy match: nothing is dropped except accents and repeated whitespace, so
    "Moyennes mensuelles" matches and "Moyenne mensuelle" does not, which is the
    difference that told recon 02 section 3.3 that the weekly note had changed
    layout.
    """
    if text is None:
        return ""
    raw = str(text).translate(_SPACES).translate(_CURRENCY)
    decomposed = unicodedata.normalize("NFKD", raw)
    stripped = "".join(ch for ch in decomposed if not unicodedata.combining(ch))
    return " ".join(stripped.split()).casefold()


#: A Drupal duplicate filename counter: one or more digits after a final
#: underscore, immediately before the extension.
_DRUPAL_COUNTER = re.compile(r"_\d+$")


def _title_of(href: str) -> str:
    """The document title a href carries, normalised and stripped of noise.

    Percent decoded, reduced to the last path segment, with the extension and any
    Drupal _N counter removed. "Historique%20du%20cours%20du%20Brent%20depuis%20
    2015%20%28moyennes%20mensuelles%29_0.xlsx" comes back as
    "historique du cours du brent depuis 2015 (moyennes mensuelles)".
    """
    path = href.split("?", 1)[0].split("#", 1)[0]
    name = unquote(path.rsplit("/", 1)[-1])
    stem = name.rsplit(".", 1)[0] if "." in name else name
    return normalise(_DRUPAL_COUNTER.sub("", stem))


# --------------------------------------------------------------------------
# Finding the files
# --------------------------------------------------------------------------

class _LinkCollector(HTMLParser):
    """Collect every href on a page, in document order.

    A parser rather than a regex for the reason worldbank.py gives: the landing
    page is a rendered government page whose markup changes with every redesign,
    and a parser fails on a shape it does not understand instead of matching
    something that merely looks like a link.
    """

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.hrefs: list[str] = []

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag.lower() != "a":
            return
        for name, value in attrs:
            if name.lower() == "href" and value:
                self.hrefs.append(value.strip())


def discover_workbook_url(
    html, title: str, *, base_url: str = LANDING_PAGE, extension: str = ".xlsx"
) -> str:
    """Find one workbook's URL on the landing page, by its title. No network.

    Args:
        html: the landing page body, text or bytes.
        title: the document title, as BRENT_DOCUMENT_TITLE or MBR_DOCUMENT_TITLE
            spell it. Compared after normalise(), with the extension and any
            Drupal _N counter stripped off the href.
        base_url: what to resolve a relative href against.
        extension: the file extension the link must carry.

    Returns:
        The absolute URL of the workbook.

    Raises:
        SourceError: when the page carries no matching link, or carries two that
            resolve differently. Both are loud on purpose, and neither falls back
            to a remembered URL: the _N counter is assigned by the CMS at upload
            time, the URL the visible label claims already returns HTTP 404, and
            an old edition that still resolved would be fetched successfully and
            cached under a vintage nobody read.
    """
    if isinstance(html, (bytes, bytearray)):
        html = bytes(html).decode("utf-8", errors="replace")

    parser = _LinkCollector()
    parser.feed(html)
    wanted = normalise(title)
    suffix = extension.lower()

    hits: list[str] = []
    for href in parser.hrefs:
        path = href.split("?", 1)[0].split("#", 1)[0]
        if not unquote(path).lower().endswith(suffix):
            continue
        if _title_of(href) != wanted:
            continue
        absolute = urljoin(base_url, href)
        if absolute not in hits:
            hits.append(absolute)

    if not hits:
        offered = sorted(
            {
                _title_of(h)
                for h in parser.hrefs
                if unquote(h.split("?", 1)[0]).lower().endswith(suffix)
            }
        )
        raise SourceError(
            "dgec: no %s link titled %r on %s. The page carried %d link(s) in "
            "total and %d %s link(s), titled: %s. The href is authoritative and "
            "the visible label is not, recon 02 section 1.2, and there is no URL "
            "to fall back to: the Drupal _N counter is assigned at upload time "
            "and the spelling the label claims returns HTTP 404. Open the page "
            "and find where the file moved to."
            % (
                suffix,
                title,
                base_url,
                len(parser.hrefs),
                len(offered),
                suffix,
                ", ".join(repr(t) for t in offered) or "none",
            )
        )
    if len(hits) > 1:
        raise SourceError(
            "dgec: %d different links titled %r on %s, %s. Picking one would be a "
            "guess about which edition is current"
            % (len(hits), title, base_url, ", ".join(hits))
        )
    return hits[0]


# --------------------------------------------------------------------------
# Parsing, no network
# --------------------------------------------------------------------------

def _open_sheet(payload, sheet: str, *, series: str):
    """Open a workbook and return (book, worksheet) for one named sheet."""
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover, an environment problem
        raise SourceError(
            "dgec %s: openpyxl is not installed and the ministry publishes an "
            "xlsx" % series
        ) from exc

    try:
        handle = (
            io.BytesIO(bytes(payload))
            if isinstance(payload, (bytes, bytearray))
            else str(payload)
        )
        book = openpyxl.load_workbook(handle, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001, openpyxl raises several types
        raise SourceError(
            "dgec %s: the download did not open as an xlsx workbook, %s: %s. That "
            "is usually an HTML error page served under the file's name"
            % (series, type(exc).__name__, exc)
        ) from exc

    if sheet not in book.sheetnames:
        raise SourceError(
            "dgec %s: the workbook has no %r sheet, sheets were %s"
            % (series, sheet, ", ".join(book.sheetnames))
        )
    return book, book[sheet]


def _rows_of(payload, sheet: str, *, series: str) -> list[list]:
    book, worksheet = _open_sheet(payload, sheet, series=series)
    try:
        return [list(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        book.close()


def _cell(row: Sequence, index: int):
    return row[index] if index < len(row) else None


def _find_row(rows: list[list], label: str, *, column: int = 0) -> int | None:
    """Index of the first row whose cell in `column` normalises to `label`."""
    for index, row in enumerate(rows):
        if normalise(_cell(row, column)) == label:
            return index
    return None


def _check_title_cell(rows: list[list], expected: str, *, series: str) -> None:
    """Refuse a workbook whose own title cell is not the one we came for.

    The title sits in B1 in both files. It is searched for over the rows above
    the monthly header rather than read from a fixed cell, because a title that
    moved down one row is not a reason to refuse a file, while a title that is
    absent altogether means this is a different table.
    """
    for row in rows[: min(6, len(rows))]:
        for cell in row:
            if normalise(cell) == expected:
                return
    seen = [
        normalise(c)
        for row in rows[: min(6, len(rows))]
        for c in row
        if normalise(c)
    ]
    raise SourceError(
        "dgec %s: no cell in the first rows of the sheet reads %r. Read: %s. The "
        "title is checked so that a workbook served under the right filename but "
        "holding a different table is refused rather than parsed"
        % (series, expected, ", ".join(repr(s) for s in seen) or "nothing")
    )


def _monthly_block(rows: list[list], *, series: str) -> tuple[int, int]:
    """The half open row range of the monthly data block, (first, stop).

    THE ANNUAL AVERAGES TRAP. Both workbooks carry an annual block below the
    monthly one keyed by a year written as text, so column A is not always a
    date, and pandas would turn '2015' into 2015-01-01 and collide it with the
    real January 2015 row. Recon 02 sections 1.3 and 1.4.

    The block is bounded rather than filtered: it starts on the row after the one
    whose column A reads "Moyennes mensuelles" and ends before the row whose
    column A reads "MOYENNES ANNUELLES". Rows inside it are then required to be
    dates, see parse. Bounding rather than filtering is what makes a layout
    change loud: if the annual heading is ever dropped or the two blocks are ever
    swapped, a text year lands inside the block and the parse stops with the cell
    named, instead of eleven annual averages quietly joining a monthly series.

    Raises:
        SourceError: when the monthly header is absent, or when the annual
            heading sits above it.
    """
    header = _find_row(rows, MONTHLY_BLOCK_LABEL)
    if header is None:
        raise SourceError(
            "dgec %s: no row whose column A reads %r, so the monthly block cannot "
            "be located. Column A of these files is not always a date, it also "
            "carries an annual averages block keyed by a text year, so the block "
            "is bounded by its headings and never guessed at"
            % (series, MONTHLY_BLOCK_LABEL)
        )

    annual = _find_row(rows, ANNUAL_BLOCK_LABEL)
    if annual is None:
        # Not an error by itself. The monthly block then runs to the end of the
        # sheet, and the date check inside parse is what protects the series: any
        # text year still present would now be inside the block and would stop
        # the parse with its cell named.
        return header + 1, len(rows)
    if annual <= header:
        raise SourceError(
            "dgec %s: the %r heading is at row %d, at or above the %r heading at "
            "row %d. The two blocks have been reordered and the monthly block can "
            "no longer be bounded safely"
            % (series, ANNUAL_BLOCK_LABEL, annual + 1, MONTHLY_BLOCK_LABEL, header + 1)
        )
    return header + 1, annual


def _unit_column(
    rows: list[list], header: int, unit: str, *, series: str, what: str
) -> int:
    """The index of the column whose header cell carries one unit. By label.

    Args:
        rows: the sheet.
        header: index of the monthly header row, the one reading "Moyennes
            mensuelles" in column A.
        unit: the normalised unit string, UNIT_USD_BBL or UNIT_EUR_T.
        what: what the column is, for the error message.

    Raises:
        SourceError: when no column carries that unit, or more than one does.
            SPEC.md section 13: a column that moves must break loudly rather than
            silently return the wrong product. The MBR sheet has two value
            columns whose only distinguishing mark is this header cell, and
            reading them by position would swap dollars per barrel for euros per
            tonne, a factor of about seven, with no other symptom.
    """
    row = rows[header] if header < len(rows) else []
    hits = [index for index in range(1, len(row)) if normalise(row[index]) == unit]
    if not hits:
        printed = [normalise(c) for c in row[1:] if normalise(c)]
        raise SourceError(
            "dgec %s: no column on the header row carries the unit %r for the %s "
            "column. The header row reads: %s. Units are matched by label and "
            "never by position, SPEC.md section 13"
            % (series, unit, what, ", ".join(repr(p) for p in printed) or "nothing")
        )
    if len(hits) > 1:
        raise SourceError(
            "dgec %s: %d columns carry the unit %r, at columns %s. Picking one "
            "would be a guess about which is the %s"
            % (series, len(hits), unit, ", ".join(str(h + 1) for h in hits), what)
        )
    return hits[0]


def _parse_month(value, *, series: str, row_number: int) -> pd.Timestamp:
    """One column A cell inside the monthly block. Anything but a date raises.

    This is the whole defence against the annual averages block. A text year
    reaching here is not coerced, not skipped and not warned about: it stops the
    parse, because a parser that silently drops rows it does not understand
    cannot tell a layout change from an empty file.
    """
    if isinstance(value, _dt.datetime):
        stamp = pd.Timestamp(value)
    elif isinstance(value, _dt.date):
        stamp = pd.Timestamp(value)
    else:
        raise SourceError(
            "dgec %s: row %d of the monthly block has %r in column A, which is "
            "not a date. Both workbooks carry an ANNUAL averages block below the "
            "monthly one, keyed by a year written as text, recon 02 sections 1.3 "
            "and 1.4, and coercing that text would turn '2015' into 2015-01-01 "
            "and collide it with the real January 2015 row. Either the block "
            "headings moved or the layout changed; look at the file"
            % (series, row_number, value)
        )
    if stamp.day != 1:
        raise SourceError(
            "dgec %s: row %d of the monthly block is dated %s, which is not the "
            "first of a month. These are monthly averages keyed by the first day "
            "of the month, so a mid month date means the sheet now means "
            "something else" % (series, row_number, stamp.date().isoformat())
        )
    return stamp.normalize()


def _parse_number(value, *, series: str, where: str) -> float:
    """One value cell. Empty becomes NaN, anything unreadable raises.

    Never a zero. SPEC.md non negotiable 1: a missing observation is NaN, and a
    cell this parser cannot read is a format change, not a margin of zero.
    """
    if value is None:
        return float("nan")
    if isinstance(value, bool):
        raise SourceError(
            "dgec %s %s: the cell holds a boolean, not a price" % (series, where)
        )
    if isinstance(value, (int, float)):
        return float(value)
    text = normalise(value)
    if text in {"", "-", "n/a", "na", "nan", "..", "..."}:
        return float("nan")
    # A French decimal comma, and a thousands separator already turned into a
    # plain space by normalise(). Rewritten deliberately rather than coerced,
    # then failed on rather than guessed at.
    candidate = text.replace(" ", "").replace(",", ".")
    try:
        return float(candidate)
    except ValueError as exc:
        raise SourceError(
            "dgec %s %s: cannot read %r as a number. If the workbook has started "
            "writing a new token where it has no value, add it deliberately "
            "rather than letting it become a number" % (series, where, value)
        ) from exc


def _parse_monthly_sheet(
    payload,
    *,
    series: str,
    sheet: str,
    title_cell: str,
    columns: Sequence[tuple[str, str, str]],
) -> pd.DataFrame:
    """Parse one of the two monthly workbooks. No network.

    Args:
        payload: the response body as bytes, or a path to a saved copy.
        series: the series name, for error messages.
        sheet: the sheet name, exact.
        title_cell: the normalised title cell to insist on.
        columns: one (output column name, normalised unit, description) triple
            per value column, in output order. Each is located by its unit label
            on the header row.

    Returns:
        A frame with a date column and the requested value columns, ascending by
        date, NaN where the workbook carried nothing.

    Raises:
        SourceError: on anything that is not this workbook. See _monthly_block,
            _unit_column and _parse_month for the three ways these files can
            change under this parser.
    """
    rows = _rows_of(payload, sheet, series=series)
    _check_title_cell(rows, title_cell, series=series)

    header = _find_row(rows, MONTHLY_BLOCK_LABEL)
    first, stop = _monthly_block(rows, series=series)
    indices = [
        _unit_column(rows, header, unit, series=series, what=what)
        for _name, unit, what in columns
    ]

    dates: list[pd.Timestamp] = []
    values: dict[str, list[float]] = {name: [] for name, _u, _w in columns}

    for index in range(first, min(stop, len(rows))):
        row = rows[index]
        if not row or all(cell is None for cell in row):
            # A blank spacer between the two blocks. Blank rows carry nothing and
            # end nothing, so they are skipped; the block boundary is what stops
            # the read, not the first empty line.
            continue
        month = _parse_month(row[0], series=series, row_number=index + 1)
        dates.append(month)
        for (name, _unit, _what), column in zip(columns, indices):
            values[name].append(
                _parse_number(
                    _cell(row, column),
                    series=series,
                    where="%s, %s" % (month.strftime("%Y-%m"), name),
                )
            )

    if not dates:
        raise SourceError(
            "dgec %s: the monthly block of sheet %r held no dated row"
            % (series, sheet)
        )

    frame = pd.DataFrame({"date": dates, **values})
    duplicates = frame.loc[frame["date"].duplicated(), "date"]
    if len(duplicates):
        raise SourceError(
            "dgec %s: duplicate month(s) in the monthly block, first is %s"
            % (series, duplicates.iloc[0].strftime("%Y-%m"))
        )
    frame = frame.sort_values("date", kind="mergesort").reset_index(drop=True)

    # Every unit cell the header row actually carries, normalised, travels with
    # the frame and ends up in the manifest. It is how a reader can check
    # UNITS_NOT_IN_FILE rather than take its word: if DGEC ever adds the
    # centimes per litre column the note prints, it appears here on the next
    # refresh instead of staying invisible because nothing asked for it.
    frame.attrs["header_units"] = [
        normalise(cell) for cell in rows[header][1:] if normalise(cell)
    ]

    first_month = frame["date"].iloc[0]
    if first_month > SERIES_START:
        raise SourceError(
            "dgec %s: the workbook starts %s, later than the documented start of "
            "%s. The file has been re cut and the history would silently shorten"
            % (series, first_month.strftime("%Y-%m"), SERIES_START.strftime("%Y-%m"))
        )
    return frame


def parse_brent_workbook(payload) -> pd.DataFrame:
    """Parse the Brent history workbook into the cache layout. No network.

    Returns:
        date, brent_usd_bbl as published, and brent_usd_t derived from it at
        crack.config.DGEC_BBL_PER_T_BRENT_NOTE, which is 7.5 bbl/t and is the
        factor the weekly note itself uses for this conversion. See the module
        docstring for why it is not 7.55.
    """
    frame = _parse_monthly_sheet(
        payload,
        series=SERIES_BRENT,
        sheet=BRENT_SHEET,
        title_cell=BRENT_TITLE_CELL,
        columns=((COLUMN_BRENT_USD_BBL, UNIT_USD_BBL, "Brent date in $/b"),),
    )
    # Derived, one multiplication, declared in the manifest under
    # derived_columns. The note prints the same quantity rounded to the whole
    # dollar, which is the form SPEC.md section 5.5's 628 anchor takes.
    frame[COLUMN_BRENT_USD_T] = frame[COLUMN_BRENT_USD_BBL] * DGEC_BBL_PER_T_BRENT_NOTE
    return frame


def parse_mbr_workbook(payload) -> pd.DataFrame:
    """Parse the MBR history workbook into the cache layout. No network.

    Returns:
        date, mbr_usd_bbl and mbr_eur_t, both as published. There is no centimes
        per litre column in this file and none is computed, see UNITS_NOT_IN_FILE.
    """
    return _parse_monthly_sheet(
        payload,
        series=SERIES_MBR,
        sheet=MBR_SHEET,
        title_cell=MBR_TITLE_CELL,
        columns=(
            (COLUMN_MBR_USD_BBL, UNIT_USD_BBL, "margin in $/b"),
            (COLUMN_MBR_EUR_T, UNIT_EUR_T, "margin in EUR/t"),
        ),
    )


# --------------------------------------------------------------------------
# Provisional
# --------------------------------------------------------------------------

def provisional_month(last_month, today=None) -> str | None:
    """Which month of this workbook, if any, DGEC would still call provisional.

    SPEC.md non negotiable 5 and section 5.3. The rule is mechanical and narrow
    on purpose:

        these workbooks publish complete months. Read on 2026-09-11 the Brent
        file ended 2026-08 and the MBR file ended 2026-07, while the weekly note
        of 4 September 2026 already carried August 2026. The provisional figure
        lives in the note, which prints '(donnees provisoires)' under the current
        month's column, and not in these files.

    So the last month here is flagged provisional only when it IS the month the
    fetch is happening in, which is the only circumstance under which DGEC would
    be publishing an incomplete month in this file. Guessing more than that would
    put a flag on a final figure, which is its own kind of lie.

    WHAT THIS RULE IS NOT. It is not an argument that a month which is over is
    final. Gate 1 self audit, finding 1.2: SPEC.md section 5.5 gives August 2026
    as 38.05 provisional on 28 August, after the month had ended, so "the month
    is over" establishes nothing by itself. Finality is established by the
    ministry dropping its own '(donnees provisoires)' marker, which happens in
    the note and is decoded by crack.sources.dgec_note into
    dgec_note_printed_monthly. This function is a narrow guard against the one
    case this workbook could get wrong on its own, and the note is the evidence.

    Args:
        last_month: the last month in the frame.
        today: the date to judge against, for tests. Defaults to today in UTC.

    Returns:
        The first day of the provisional month as yyyy-mm-dd, or None.
    """
    stamp = pd.Timestamp(last_month)
    now = pd.Timestamp(today) if today is not None else pd.Timestamp.utcnow()
    if (stamp.year, stamp.month) == (now.year, now.month):
        return stamp.normalize().strftime("%Y-%m-%d")
    return None


# --------------------------------------------------------------------------
# Fetching
# --------------------------------------------------------------------------

#: The landing page, kept for LANDING_CACHE_SECONDS so that running both
#: adapters in one process asks the ministry for the same HTML once. Holds
#: (monotonic seconds, body bytes).
_LANDING_CACHE: list = []


def clear_landing_cache() -> None:
    """Forget the cached landing page. Called by tests, and safe at any time."""
    _LANDING_CACHE.clear()


def fetch_landing_page(*, delay: float = 1.0, force: bool = False) -> bytes:
    """The landing page body, fetched at most once per LANDING_CACHE_SECONDS."""
    if not force and _LANDING_CACHE:
        stamped, body = _LANDING_CACHE[0]
        if time.monotonic() - stamped < LANDING_CACHE_SECONDS:
            return body
    response = http_get(LANDING_PAGE, delay=delay, retries=3, timeout=45)
    body = response.content
    if not body:
        raise SourceError("dgec: the landing page %s came back empty" % LANDING_PAGE)
    _LANDING_CACHE.clear()
    _LANDING_CACHE.append((time.monotonic(), body))
    return body


_XLSX_ACCEPT = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*;q=0.8"
)


def _fetch_workbook(title: str, *, series: str, delay: float) -> tuple[bytes, str, str]:
    """Discover one workbook's URL, fetch it, and report what the server said.

    Returns:
        (body, url, last_modified). last_modified is the HTTP header verbatim, or
        an empty string when the server sends none. It is the only vintage stamp
        these files carry, since neither workbook holds an "updated on" cell the
        way the World Bank pink sheet does.
    """
    html = fetch_landing_page(delay=delay)
    url = discover_workbook_url(html, title, base_url=LANDING_PAGE)
    response = http_get(
        url, headers={"Accept": _XLSX_ACCEPT}, delay=delay, retries=3, timeout=90
    )
    body = response.content
    if not body:
        raise SourceError("dgec %s: the workbook download from %s was empty" % (series, url))
    return body, url, response.headers.get("Last-Modified", "") or ""


def fetch_brent(*, delay: float = 1.0) -> tuple[pd.DataFrame, str, str]:
    """Discover, fetch and parse the Brent history workbook."""
    body, url, last_modified = _fetch_workbook(
        BRENT_DOCUMENT_TITLE, series=SERIES_BRENT, delay=delay
    )
    return parse_brent_workbook(body), url, last_modified


def fetch_mbr(*, delay: float = 1.0) -> tuple[pd.DataFrame, str, str]:
    """Discover, fetch and parse the MBR history workbook."""
    body, url, last_modified = _fetch_workbook(
        MBR_DOCUMENT_TITLE, series=SERIES_MBR, delay=delay
    )
    return parse_mbr_workbook(body), url, last_modified


# --------------------------------------------------------------------------
# The adapters
# --------------------------------------------------------------------------

class _DgecWorkbookAdapter(Adapter):
    """What the two workbook adapters share: discovery, vintage, provisional."""

    source = "DGEC via ecologie.gouv.fr, assessments credited to Reuters"
    page_url = LANDING_PAGE
    url = LANDING_PAGE
    frequency = "monthly"
    method = "published"
    committable = True
    date_col = "date"

    #: The document title to look for on the landing page.
    document_title: str = ""

    def __init__(self, *, delay: float = 1.0, today=None):
        """
        Args:
            delay: polite pause before each request, seconds. Both adapters share
                one landing page fetch, see fetch_landing_page.
            today: the date provisional_month() judges against. Tests pass it so
                they give the same answer next year.
        """
        self.delay = float(delay)
        self.today = today
        #: the URL discovered on the landing page, set by fetch()
        self.discovered_url: str | None = None
        #: the HTTP Last-Modified header of the workbook, set by fetch()
        self.last_modified: str = ""
        #: every unit cell the header row carried, set by fetch(). See the
        #: comment in _parse_monthly_sheet.
        self.header_units: list[str] = []

    def _stamp(self, frame: pd.DataFrame, url: str, last_modified: str) -> None:
        """Record the URL, the vintage and the provisional flag from one fetch."""
        self.discovered_url = url
        self.url = url
        self.last_modified = last_modified
        self.header_units = list(frame.attrs.get("header_units", []))
        last_month = pd.Timestamp(frame["date"].max())
        filename = unquote(url.rsplit("/", 1)[-1])
        self.vintage = "%s, last month %s%s" % (
            filename,
            last_month.strftime("%Y-%m"),
            ", Last-Modified %s" % last_modified if last_modified else "",
        )
        self.provisional_from = provisional_month(last_month, self.today)

    #: The two openings _provisional_sentence can produce. The sentence is always
    #: the LAST thing in either adapter's note, which is what lets offline_note
    #: restate it without knowing anything only a fetch could know.
    PROVISIONAL_OPENINGS = ("No provisional month", "PROVISIONAL:")

    def offline_note(self, previous_note: str, frame: pd.DataFrame) -> str:
        """Restate the provisional sentence of a carried forward note. No fetch.

        scripts/refresh.py --offline keeps the note the last real run wrote,
        because that note names the URL the run discovered and the run's own
        measurements, none of which offline can know. But the provisional
        sentence is not one of those things: it depends only on the last month in
        the committed cache and on today's date, both of which offline has.

        Without this hook a correction to the wording of that sentence would sit
        in the source and never reach the manifest until somebody happened to run
        an online fetch, and the manifest is what the provenance panel prints.
        Gate 1 self audit, finding 1.2, is exactly such a correction: the sentence
        used to argue from "the month is already over", which does not establish
        finality.

        Returns the previous note unchanged when it carries no provisional
        sentence to replace, so this can never invent one.
        """
        if not previous_note:
            return previous_note
        self.provisional_from = provisional_month(
            pd.Timestamp(frame[self.date_col].max()), self.today
        )
        cuts = [
            previous_note.find(opening)
            for opening in self.PROVISIONAL_OPENINGS
            if previous_note.find(opening) >= 0
        ]
        if not cuts:
            return previous_note
        return previous_note[: min(cuts)] + self._provisional_sentence(frame)

    def _provisional_sentence(self, frame: pd.DataFrame) -> str:
        last_month = pd.Timestamp(frame["date"].max()).strftime("%Y-%m")
        if self.provisional_from:
            return (
                "PROVISIONAL: %s is the month this fetch happened in, so it is "
                "incomplete and DGEC will revise it. Recon 02 section 3.5 "
                "measured a 2 $/b revision on August 2024 between one note and "
                "the next. Never blend it into a final series." % last_month
            )
        # THE REASON HERE USED TO BE THE WRONG REASON, AND THE GATE 1 SELF AUDIT
        # WAS RIGHT TO SAY SO, finding 1.2. It said the last month is not
        # provisional because it "is already over", and a month being over is not
        # the same as its figure being final: SPEC.md section 5.5 gives August
        # 2026 as 38.05 PROVISIONAL on 28 August, which is after the month ended.
        # The conclusion is right and this is why. Finality is established from
        # the note, not from the calendar.
        return (
            "No provisional month, and the reason is the note rather than the "
            "calendar. This workbook publishes complete months and its last is "
            "%s. A month being over does NOT make its figure final: SPEC.md "
            "section 5.5 gives August 2026 as 38.05 provisional on 28 August, "
            "after the month had ended. What establishes finality is the "
            "ministry's own marker: DGEC prints '(donnees provisoires)' under the "
            "provisional column of the weekly note and drops it when the figure "
            "goes final, and the note of 4 September 2026 prints 38,05 for August "
            "2026 in a column with no marker. That is decoded by "
            "crack.sources.dgec_note, carried per row in "
            "dgec_note_printed_monthly, and asserted by "
            "tests/test_dgec_anchor.py::"
            "test_august_2026_is_final_and_is_not_flagged_provisional. WHAT "
            "REMAINS UNCERTAIN: this workbook carries no provisional marker of "
            "its own at all, so for any month where no preserved note survives "
            "there is no positive evidence of finality, only the absence of a "
            "reason to doubt it. The corpus is ten notes." % last_month
        )

    def _entry(self, *, status: str, frame, note: str) -> dict:
        entry = super()._entry(status=status, frame=frame, note=note)
        entry["licence"] = "Licence Ouverte 2.0 (Etalab)"
        entry["attribution"] = (
            "DGEC, Direction generale de l'energie et du climat, ministere de la "
            "Transition ecologique. Assessments credited 'Source : DGEC-REUTERS'. "
            "Licence Ouverte 2.0 requires the source and the date of last update "
            "of the data reused, which is the vintage field."
        )
        if self.discovered_url:
            entry["discovered_url"] = self.discovered_url
        if self.last_modified:
            entry["last_modified"] = self.last_modified
        if self.header_units:
            entry["header_units"] = list(self.header_units)
        return entry


class DgecBrentMonthly(_DgecWorkbookAdapter):
    """Brent date, monthly averages, as DGEC publishes them.

    The page states, verbatim: "Les moyennes mensuelles du fichier Excel
    ci-dessous sont les moyennes des cours quotidiens du Brent date en cloture a
    Londres. Les moyennes annuelles sont les moyennes des moyennes mensuelles."
    So this is a mean of daily London closes assessed by Reuters, which is a
    DIFFERENT assessment from EIA Europe Brent spot, and that is what makes the
    SPEC.md section 5.5 cross check against FRED DCOILBRENTEU a real test rather
    than a tautology.
    """

    name = SERIES_BRENT
    document_title = BRENT_DOCUMENT_TITLE
    unit = "USD per barrel, plus USD per tonne derived at 7.5 bbl/t"
    licence_note = _registered_brent.licence_note
    required_cols = ("date", COLUMN_BRENT_USD_BBL, COLUMN_BRENT_USD_T)
    bounds = {
        COLUMN_BRENT_USD_BBL: BOUNDS_BRENT_USD_BBL,
        COLUMN_BRENT_USD_T: BOUNDS_PRODUCT_USD_T,
    }
    min_rows = MIN_ROWS_BRENT
    min_observations = {
        COLUMN_BRENT_USD_BBL: MIN_ROWS_BRENT,
        COLUMN_BRENT_USD_T: MIN_ROWS_BRENT,
    }
    observation_column = COLUMN_BRENT_USD_BBL

    def fetch(self) -> pd.DataFrame:
        frame, url, last_modified = fetch_brent(delay=self.delay)
        self._stamp(frame, url, last_modified)
        self.note = (
            "Brent date monthly averages, means of the daily London closes, read "
            "from sheet %r of the ministry's own history workbook. %d months, %s "
            "to %s. The $/b column is DGEC's; %s is DERIVED here by multiplying "
            "it by %g bbl/t, which is the factor the weekly note itself uses to "
            "print Brent date in its $/t quotation table, recovered by recon 02 "
            "section 4.3 on five month pairs and reproducing SPEC.md section 5.5's "
            "July 2026 anchor of 628 $/t exactly. It is NOT the %g bbl/t the "
            "methodology note states for the CAF Brent inside the margin "
            "calculation; the two are used in different places and swapping them "
            "is a silent 0.67 percent error on the crude leg. The link was "
            "discovered from the landing page because the Drupal _N counter in "
            "the path is unstable and the spelling the visible label claims "
            "returns HTTP 404, recon 02 section 1.2: this run used %s. %s"
            % (
                BRENT_SHEET,
                len(frame),
                frame["date"].min().strftime("%Y-%m"),
                frame["date"].max().strftime("%Y-%m"),
                COLUMN_BRENT_USD_T,
                DGEC_BBL_PER_T_BRENT_NOTE,
                DGEC_BBL_PER_T_BRENT_MARGIN,
                url,
                self._provisional_sentence(frame),
            )
        )
        return frame

    def _entry(self, *, status: str, frame, note: str) -> dict:
        entry = super()._entry(status=status, frame=frame, note=note)
        entry["derived_columns"] = [
            {
                "column": COLUMN_BRENT_USD_T,
                "from": COLUMN_BRENT_USD_BBL,
                "factor_bbl_per_t": DGEC_BBL_PER_T_BRENT_NOTE,
                "note": (
                    "Derived by this project, not published in this workbook. The "
                    "weekly note prints the same quantity rounded to the whole "
                    "dollar. The factor is the note's own, recovered by recon 02 "
                    "section 4.3 on five month pairs, and must never be confused "
                    "with the %g bbl/t the methodology note states for the CAF "
                    "Brent inside the margin calculation."
                    % DGEC_BBL_PER_T_BRENT_MARGIN
                ),
            }
        ]
        entry["source_line"] = "Source : Reuters / DGEC"
        entry["assessment_note"] = (
            "A mean of daily Brent date London closes assessed by Reuters. This is "
            "a different assessment from EIA Europe Brent spot, which is what the "
            "SPEC.md section 5.5 cross check compares it against."
        )
        return entry


class DgecMbrMonthly(_DgecWorkbookAdapter):
    """The official gross refining margin on Brent, monthly, $/b and EUR/t.

    SPEC.md section 4.3 layer 1 calls this the headline NWE margin series, and
    section 5.5 anchors eight of its values. Every one of those eight reproduces
    from this file to two decimal places, recon 02 section 4.1, and
    tests/test_dgec_anchor.py asserts them against the committed cache.
    """

    name = SERIES_MBR
    document_title = MBR_DOCUMENT_TITLE
    unit = "USD per barrel and EUR per tonne"
    licence_note = _registered_mbr.licence_note
    required_cols = ("date", COLUMN_MBR_USD_BBL, COLUMN_MBR_EUR_T)
    bounds = {
        COLUMN_MBR_USD_BBL: BOUNDS_MARGIN_USD_BBL,
        COLUMN_MBR_EUR_T: BOUNDS_MBR_EUR_T,
    }
    min_rows = MIN_ROWS_MBR
    min_observations = {
        COLUMN_MBR_USD_BBL: MIN_ROWS_MBR,
        COLUMN_MBR_EUR_T: MIN_ROWS_MBR,
    }
    observation_column = COLUMN_MBR_USD_BBL

    def fetch(self) -> pd.DataFrame:
        frame, url, last_modified = fetch_mbr(delay=self.delay)
        self._stamp(frame, url, last_modified)
        self.note = (
            "Gross refining margin on Brent, the marge brute de raffinage sur "
            "Brent, read from sheet %r of the ministry's own history workbook. %d "
            "months, %s to %s, in $/b and EUR/t. Both value columns are located "
            "by the unit label on the header row and never by position: they are "
            "adjacent, they differ by a factor of about seven, and swapping them "
            "would have no other symptom. There is no centimes per litre column "
            "in this file and none is computed here. This is DGEC's own published "
            "figure and it is never overwritten by a replication. The link was "
            "discovered from the landing page, this run used %s. %s"
            % (
                MBR_SHEET,
                len(frame),
                frame["date"].min().strftime("%Y-%m"),
                frame["date"].max().strftime("%Y-%m"),
                url,
                self._provisional_sentence(frame),
            )
        )
        return frame

    def _entry(self, *, status: str, frame, note: str) -> dict:
        entry = super()._entry(status=status, frame=frame, note=note)
        entry["method_change"] = dict(METHOD_CHANGE)
        entry["note_leads_file"] = NOTE_LEADS_FILE
        entry["units_not_in_file"] = UNITS_NOT_IN_FILE
        entry["replication_note"] = (
            "SPEC.md section 4.3 layer 2 asks for this margin to be recomputed "
            "from DGEC's own published quotations to within 0.50 $/bbl. Recon 02 "
            "section 5.6 established that it cannot be: EuroBOB, essence export, "
            "naphta, propane and butane, the PEG Nord day ahead gas price, the "
            "GRTgaz transport tariff and the Sullom Voe to Le Havre Aframax "
            "freight are never published, which is five of the ten product "
            "quotations plus the gas and freight legs. The spec's own fallback "
            "applies: name the missing inputs and use the official series. The "
            "partial attribution of section 4.3 layer 3 over the five published "
            "quotations, 59.3 percent of the mass yield, is what can honestly be "
            "built, with everything else on a residual line."
        )
        return entry


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------

def _report(entry: dict) -> None:
    print(
        "%-22s rows=%-5d %s to %s gaps=%d status=%s provisional_from=%s"
        % (
            entry["series"],
            entry["rows"],
            entry["first_date"],
            entry["last_date"],
            len(entry["gaps"]),
            entry["status"],
            entry["provisional_from"],
        )
    )
    print("%-22s vintage=%s" % ("", entry["vintage"]))


def main(argv: Iterable[str] | None = None) -> int:
    """Run both adapters and print an ASCII report. Returns an exit code."""
    failures = 0
    for adapter in (DgecBrentMonthly(), DgecMbrMonthly()):
        try:
            entry = adapter.run()
        except Exception as exc:  # noqa: BLE001, the CLI reports and exits non zero
            print("%s FAILED: %s: %s" % (adapter.name, type(exc).__name__, exc))
            failures += 1
            continue
        _report(entry)
    return 1 if failures else 0


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(main())
